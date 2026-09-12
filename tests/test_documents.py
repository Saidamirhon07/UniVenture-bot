import asyncio
import base64
import io
import json
import logging
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from docx import Document
from backend.documents import DocumentError, SubmissionDocument, prepare_document, attach_document, current_document, MAX_BYTES
from backend import legacy
from backend.prompts import compact_evaluation_messages


def pdf_fixture(pages=1):
    # Minimal valid PDF, no extra test dependency required.
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>']
    kids = ' '.join(f'{3+i*2} 0 R' for i in range(pages))
    objects.append(f'<< /Type /Pages /Kids [{kids}] /Count {pages} >>'.encode())
    for i in range(pages):
        objects.extend([f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents {4+i*2} 0 R >>'.encode(), b'<< /Length 0 >>\nstream\n\nendstream'])
    output = b'%PDF-1.4\n'; offsets = [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(output)); output += f'{n} 0 obj\n'.encode()+obj+b'\nendobj\n'
    xref=len(output); output += f'xref\n0 {len(objects)+1}\n0000000000 65535 f \n'.encode()
    output += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    output += f'trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    return output

class DocumentTests(unittest.TestCase):
    def test_pdf_original_bytes_preserved_and_not_in_metadata(self):
        data = pdf_fixture()
        document = prepare_document('../../student.pdf', data)
        self.assertEqual(document.name, 'student.pdf')
        self.assertEqual(document.pages, 1)
        self.assertEqual(base64.b64decode(document.pdf_data.split(',')[1]),data)
        self.assertNotIn('pdf_data', document.source)
        self.assertTrue(document.source['requires_reupload'])

    def test_docx_tables_and_paragraphs_keep_order(self):
        document = Document();document.add_paragraph('Before')
        document.add_table(rows=1,cols=2).cell(0,0).text='Evidence in table'
        document.add_paragraph('After');stream=io.BytesIO();document.save(stream)
        result=prepare_document('activity.docx',stream.getvalue())
        self.assertLess(result.text.index('Before'),result.text.index('Evidence in table'))
        self.assertLess(result.text.index('Evidence in table'),result.text.index('After'))

    def test_unsupported_invalid_empty_large_and_too_many_pages(self):
        for name,data in [('bad.pdf',b'not pdf'),('bad.docx',b'not zip'),('bad.exe',b'x'),('empty.txt',b''),('long.txt',b'a'*30001),('large.txt',b'a'*(MAX_BYTES+1)),('bad.txt',b'\xff'),('many.pdf',pdf_fixture(21))]:
            with self.subTest(name=name), self.assertRaises(DocumentError):prepare_document(name,data)

    def test_docx_expansion_limit(self):
        import zipfile
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml',b'0'*(21*1024*1024))
        with self.assertRaisesRegex(DocumentError,'expands'):prepare_document('bomb.docx',stream.getvalue())

    def test_rubric_prompts_include_evidence_and_nullable_scores(self):
        from backend.prompts import EVALUATION_SPECS
        for topic in EVALUATION_SPECS:
            messages=compact_evaluation_messages(topic,'student work','','')
            self.assertIn('"criteria"',messages[0]['content'])
            self.assertIn('null if evidence is insufficient',messages[0]['content'])
            self.assertIn('never assess pronunciation',messages[0]['content'])

class DocumentIsolationTests(unittest.IsolatedAsyncioTestCase):
    async def test_concurrent_requests_keep_their_own_file(self):
        async def request(name):
            token=current_document.set(SubmissionDocument(name,'',name))
            try:
                await asyncio.sleep(0)
                attached=attach_document([{'role':'system','content':'JSON'}])
                self.assertEqual(attached[-1]['content'][0]['file']['filename'],name)
            finally:current_document.reset(token)
        await asyncio.gather(request('a.pdf'),request('b.pdf'))
        self.assertIsNone(current_document.get())
        original=[{'role':'user','content':'typed'}]
        self.assertEqual(attach_document(original),original)

    async def test_json_retry_keeps_file_without_mutating_original_messages(self):
        bot=SimpleNamespace(STRONG_MODEL='gpt-4.1',FAST_MODEL='gpt-4o-mini',logging=logging,
            openai_chat=AsyncMock(side_effect=['invalid','{"headline":"review"}']))
        document=prepare_document('student.pdf',pdf_fixture())
        token=current_document.set(document)
        messages=[{'role':'system','content':'Return JSON'}]
        try:
            with patch.object(legacy,'module',return_value=bot):
                self.assertEqual(json.loads(await legacy.ask_ai(messages))['headline'],'review')
        finally:current_document.reset(token)
        self.assertEqual(len(messages),1)
        for call in bot.openai_chat.call_args_list:
            self.assertTrue(any(isinstance(m['content'],list) and m['content'][0]['type']=='file' for m in call.args[1]))

if __name__=='__main__':unittest.main()
